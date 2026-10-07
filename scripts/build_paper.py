"""Build the research paper (DOCX, IEEE-style two-column layout) from the analysis results.

Every number in the text is read from results/results.json or results/tables/*.csv,
so the paper always matches the executed notebook.

Usage (from the repository root):
    python scripts/build_paper.py      # writes report/Student_Depression_Research_Paper.docx
    python scripts/export_pdf.py       # refreshes the contents pages in MS Word and exports the PDFs
"""
import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "results"
TAB = RES / "tables"
FIG = RES / "figures"
LOGO = ROOT / "assets" / "iitkgp_logo.png"
OUT = ROOT / "report" / "Student_Depression_Research_Paper.docx"

R = json.loads((RES / "results.json").read_text())
T = {p.stem: pd.read_csv(p) for p in TAB.glob("*.csv")}
QS_FRACTION = re.search(r"A fraction ([0-9.]+) of observations",
                        (RES / "model_output" / "logit_health_summary.txt").read_text()).group(1)

TITLE = "Determinants of Depression among University Students: An Econometric Analysis using Logit and Probit Models"
AUTHORS = ["Atharv Lodhi", "Sourav Choudhary", "Piyush Bisht"]
FONT = "Times New Roman"
MID, HIGH = "Mid living expense", "High living expense"
CRAMER = "Cramer's V"

PAGE_W, MARGIN, GAP = 21.0, 1.9, 0.6
COL_W = (PAGE_W - 2 * MARGIN - GAP) / 2      # width of one column (cm)
FULL_W = PAGE_W - 2 * MARGIN                 # width spanning both columns (cm)

# ----------------------------------------------------------------------------- references (IEEE, numbered by first citation)
REFERENCES = {
    "ibrahim": 'A. K. Ibrahim, S. J. Kelly, C. E. Adams, and C. Glazebrook, "A systematic review of studies of depression '
               'prevalence in university students," J. Psychiatric Res., vol. 47, no. 3, pp. 391–400, 2013.',
    "hysen": 'A. Hysenbegasi, S. L. Hass, and C. R. Rowland, "The impact of depression on the academic productivity of '
             'university students," J. Mental Health Policy Econ., vol. 8, no. 3, pp. 145–151, 2005.',
    "eisenberg": 'D. Eisenberg, S. E. Gollust, E. Golberstein, and J. L. Hefner, "Prevalence and correlates of depression, '
                 'anxiety, and suicidality among university students," Amer. J. Orthopsychiatry, vol. 77, no. 4, '
                 'pp. 534–542, 2007.',
    "monroe": 'S. M. Monroe and A. D. Simons, "Diathesis–stress theories in the context of life stress research: '
              'Implications for the depressive disorders," Psychol. Bull., vol. 110, no. 3, pp. 406–425, 1991.',
    "cohen": 'S. Cohen and T. A. Wills, "Stress, social support, and the buffering hypothesis," Psychol. Bull., vol. 98, '
             'no. 2, pp. 310–357, 1985.',
    "lorant": 'V. Lorant, D. Deliège, W. Eaton, A. Robert, P. Philippot, and M. Ansseau, "Socioeconomic inequalities in '
              'depression: A meta-analysis," Amer. J. Epidemiol., vol. 157, no. 2, pp. 98–112, 2003.',
    "nolen": 'S. Nolen-Hoeksema, "Sex differences in unipolar depression: Evidence and theory," Psychol. Bull., vol. 101, '
             'no. 2, pp. 259–282, 1987.',
    "beck": "A. T. Beck, Depression: Clinical, Experimental, and Theoretical Aspects. New York, NY, USA: Harper & Row, 1967.",
    "kawachi": 'I. Kawachi and L. F. Berkman, "Social ties and mental health," J. Urban Health, vol. 78, no. 3, '
               'pp. 458–467, 2001.',
    "richardson": 'T. Richardson, P. Elliott, and R. Roberts, "The relationship between personal unsecured debt and mental '
                  'and physical health: A systematic review and meta-analysis," Clin. Psychol. Rev., vol. 33, no. 8, '
                  'pp. 1148–1162, 2013.',
    "whitton": 'S. W. Whitton and M. A. Whisman, "Relationship satisfaction instability and depression," J. Family Psychol., '
               'vol. 24, no. 6, pp. 791–794, 2010.',
    "bulloch": 'A. G. Bulloch, J. V. Williams, D. H. Lavorato, and S. B. Patten, "The relationship between major depression '
               'and marital disruption is bidirectional," Depression Anxiety, vol. 26, no. 12, pp. 1172–1177, 2009.',
    "kessler": 'R. C. Kessler, "Epidemiology of women and depression," J. Affective Disorders, vol. 74, no. 1, pp. 5–13, 2003.',
    "moussavi": 'S. Moussavi, S. Chatterji, E. Verdes, A. Tandon, V. Patel, and B. Ustun, "Depression, chronic diseases, '
                'and decrements in health: Results from the World Health Surveys," Lancet, vol. 370, no. 9590, '
                'pp. 851–858, 2007.',
    "wooldridge": "J. M. Wooldridge, Econometric Analysis of Cross Section and Panel Data, 2nd ed. Cambridge, MA, USA: "
                  "MIT Press, 2010.",
    "long": "J. S. Long and J. Freese, Regression Models for Categorical Dependent Variables Using Stata, 3rd ed. "
            "College Station, TX, USA: Stata Press, 2014.",
    "hosmer": "D. W. Hosmer, S. Lemeshow, and R. X. Sturdivant, Applied Logistic Regression, 3rd ed. Hoboken, NJ, USA: "
              "Wiley, 2013.",
    "seabold": 'S. Seabold and J. Perktold, "Statsmodels: Econometric and statistical modeling with Python," in Proc. 9th '
               'Python in Science Conf., 2010, pp. 92–96.',
    "stata": "StataCorp, Stata Statistical Software: Release 17. College Station, TX, USA: StataCorp LLC, 2021.",
}
CITED = []


def cite_numbers(text):
    """Replace [@key] or [@a; @b] with IEEE numbers [1], [2] in order of first citation."""
    def repl(m):
        nums = []
        for key in re.findall(r"@(\w+)", m.group(0)):
            if key not in CITED:
                CITED.append(key)
            nums.append(CITED.index(key) + 1)
        return ", ".join(f"[{n}]" for n in sorted(nums))
    return re.sub(r"\[@\w+(?:;\s*@\w+)*\]", repl, text)


# ----------------------------------------------------------------------------- result helpers
def row(table, label, col="Variable"):
    return T[table].set_index(col).loc[label]


def ame(label, model="logit"):
    return row(f"{model}_marginal_effects_ame", label)


def coef(label, model="logit"):
    return row(f"{model}_coefficients", label)


def odds(label):
    return row("logit_odds_ratios", label)


def rate(characteristic, group):
    t = T["depression_rates_by_group"]
    return t[(t["Characteristic"] == characteristic) & (t["Group"] == group)].iloc[0]


def fmt_p(p):
    return "p < 0.001" if p < 0.001 else f"p = {p:.3f}"


def pp(x):
    return f"{100 * x:+.1f} percentage points"


# ----------------------------------------------------------------------------- document + styles
doc = Document()
COUNTERS = {"fig": 0, "tab": 0, "sec": 0, "sub": 0}


def roman(n):
    out = ""
    for v, s in [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]:
        while n >= v:
            out, n = out + s, n - v
    return out


def set_font(style, size, bold=None, italic=None):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), FONT)


def setup_styles():
    st = doc.styles
    set_font(st["Normal"], 10)
    pf = st["Normal"].paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(0), Pt(0), 1.0
    pf.first_line_indent = Cm(0.35)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    h1 = st["Heading 1"]
    set_font(h1, 10, bold=False)
    h1.font.small_caps = True
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before, h1.paragraph_format.space_after = Pt(10), Pt(5)
    h1.paragraph_format.first_line_indent = Cm(0)
    h1.paragraph_format.keep_with_next = True

    h2 = st["Heading 2"]
    set_font(h2, 10, bold=False, italic=True)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before, h2.paragraph_format.space_after = Pt(6), Pt(3)
    h2.paragraph_format.first_line_indent = Cm(0)
    h2.paragraph_format.keep_with_next = True

    for name, align in [("Figure Caption", WD_ALIGN_PARAGRAPH.JUSTIFY), ("Table Caption", WD_ALIGN_PARAGRAPH.CENTER)]:
        s = st.add_style(name, 1)
        s.base_style = st["Normal"]
        set_font(s, 8)
        s.paragraph_format.alignment = align
        s.paragraph_format.first_line_indent = Cm(0)
        s.paragraph_format.space_before = Pt(3)
        s.paragraph_format.space_after = Pt(8)

    code = st.add_style("Code Block", 1)
    code.base_style = st["Normal"]
    set_font(code, 7)
    code.font.name = "Consolas"
    code.element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code.element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.paragraph_format.first_line_indent = Cm(0)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    eq = st.add_style("Equation", 1)
    eq.base_style = st["Normal"]
    eq.font.italic = True
    eq.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eq.paragraph_format.first_line_indent = Cm(0)
    eq.paragraph_format.space_before = Pt(4)
    eq.paragraph_format.space_after = Pt(4)

    lb = st["List Bullet"]
    set_font(lb, 10)
    lb.paragraph_format.first_line_indent = Cm(0)


def add_rich(p, text, size=None, bold=False, italic=False):
    """Adds text with **bold** spans and [@ref] citations."""
    text = cite_numbers(text)
    for i, part in enumerate(text.split("**")):
        if part:
            run = p.add_run(part)
            run.bold = bold or i % 2 == 1
            run.italic = italic
            if size:
                run.font.size = Pt(size)
    return p


def para(text="", indent=True, align=None, size=None, bold=False, italic=False, style=None, space_after=None):
    p = doc.add_paragraph(style=style)
    if text:
        add_rich(p, text, size=size, bold=bold, italic=italic)
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        add_rich(p, item)
        p.paragraph_format.space_after = Pt(1)


def heading(title):
    COUNTERS["sec"] += 1
    COUNTERS["sub"] = 0
    return doc.add_heading(f"{roman(COUNTERS['sec'])}. {title}", level=1)


def subheading(title):
    COUNTERS["sub"] += 1
    return doc.add_heading(f"{chr(64 + COUNTERS['sub'])}. {title}", level=2)


def fig_ref(offset=1):
    return f"Fig. {COUNTERS['fig'] + offset}"


def tab_ref(offset=1):
    return f"Table {roman(COUNTERS['tab'] + offset)}"


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    for kind, text in (("begin", None), ("instr", instruction), ("separate", None), ("text", "Update field."), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


# ----------------------------------------------------------------------------- sections / columns
def set_columns(section, num):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(Cm(GAP).twips)))


def drop_page_restart(section):
    pg = section._sectPr.find(qn("w:pgNumType"))
    if pg is not None and pg.get(qn("w:start")) is not None:
        del pg.attrib[qn("w:start")]


def new_section(cols, new_page=False):
    sec = doc.add_section(WD_SECTION.NEW_PAGE if new_page else WD_SECTION.CONTINUOUS)
    set_columns(sec, cols)
    drop_page_restart(sec)
    return sec


class Wide:
    """Content inside spans both columns (like IEEE figure*/table*)."""

    def __enter__(self):
        new_section(1)

    def __exit__(self, *exc):
        new_section(2)


def figure(name, text, width_cm=None, wide=False):
    width = width_cm or (FULL_W * 0.8 if wide else COL_W)

    def body():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(FIG / f"{name}.png"), width=Cm(width))
        COUNTERS["fig"] += 1
        cap = doc.add_paragraph(style="Figure Caption")
        cap.add_run(f"Fig. {COUNTERS['fig']}. {text}")

    if wide:
        with Wide():
            body()
    else:
        body()


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, sz in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), "000000")
        borders.append(el)


def table(df, text, wide=False, font_size=7.5, decimals=3, col_widths=None):
    def body():
        COUNTERS["tab"] += 1
        cap = doc.add_paragraph(style="Table Caption")
        cap.paragraph_format.space_after = Pt(3)
        cap.paragraph_format.keep_with_next = True
        r1 = cap.add_run(f"Table {roman(COUNTERS['tab'])}")
        r1.font.small_caps = True
        cap.add_run().add_break()
        r2 = cap.add_run(text)
        r2.font.small_caps = True

        frame = df.copy()
        for c in frame.columns:
            if pd.api.types.is_float_dtype(frame[c]):
                is_p = any(k in c for k in ("P>|z|", "p-value", " p", "Prob >"))
                frame[c] = frame[c].map(lambda v: "" if pd.isna(v) else
                                        ("<0.001" if is_p and v < 0.001 else f"{v:.{decimals}f}"))
        frame = frame.fillna("").astype(str).replace("nan", "")
        t = doc.add_table(rows=1, cols=len(frame.columns))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, c in enumerate(frame.columns):
            run = t.rows[0].cells[j].paragraphs[0].add_run(str(c))
            run.bold = True
            run.font.size = Pt(font_size)
        for _, r in frame.iterrows():
            cells = t.add_row().cells
            for j, v in enumerate(r):
                cells[j].paragraphs[0].add_run(v).font.size = Pt(font_size)
        for i, tr in enumerate(t.rows):
            tr._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for j, cell in enumerate(tr.cells):
                pf = cell.paragraphs[0].paragraph_format
                pf.first_line_indent = Cm(0)
                pf.space_after = pf.space_before = Pt(0.5)
                pf.keep_with_next = i < len(t.rows) - 1
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                edges = {}
                if i == 0:
                    edges = {"top": 12, "bottom": 6}
                if i == len(t.rows) - 1:
                    edges["bottom"] = 12
                if edges:
                    set_cell_border(cell, **edges)
        tbl_pr = t._tbl.tblPr
        margins = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), str(int(Cm(0.06).twips)))
            m.set(qn("w:type"), "dxa")
            margins.append(m)
        tbl_pr.append(margins)
        if col_widths:
            t.autofit = False
            layout = OxmlElement("w:tblLayout")
            layout.set(qn("w:type"), "fixed")
            tbl_pr.append(layout)
            for j, (gc, w) in enumerate(zip(t._tbl.tblGrid.findall(qn("w:gridCol")), col_widths)):
                gc.set(qn("w:w"), str(int(Cm(w).twips)))
                for tr in t.rows:
                    tr.cells[j].width = Cm(w)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.first_line_indent = Cm(0)
        spacer.paragraph_format.space_after = Pt(4)

    if wide:
        with Wide():
            body()
    else:
        body()


def note(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text).font.size = Pt(7.5)
    return p


def code_block(text):
    for line in text.rstrip().splitlines():
        p = doc.add_paragraph(style="Code Block")
        p.add_run(line if line.strip() else " ")
    doc.add_paragraph().paragraph_format.first_line_indent = Cm(0)


def footer_page_number(sect):
    sect.footer.is_linked_to_previous = False
    p = sect.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    add_field(p, "PAGE")


def cover_page():
    """Front page in the format of the course reports: title, submitted to, team, IIT KGP logo, department."""
    C = WD_ALIGN_PARAGRAPH.CENTER
    para("", indent=False, space_after=60)
    para(TITLE, indent=False, align=C, size=17, bold=True, space_after=34)
    para("Report submitted to", indent=False, align=C, size=12, space_after=2)
    para("Dr. Prantik Bagchi", indent=False, align=C, size=13, bold=True, space_after=1)
    para("Department of Humanities and Social Sciences", indent=False, align=C, size=12, space_after=1)
    para("Indian Institute of Technology Kharagpur", indent=False, align=C, size=12, space_after=40)
    para("Submitted by", indent=False, align=C, size=12, bold=True, space_after=8)
    for name in AUTHORS:
        para(name, indent=False, align=C, size=12, space_after=2)
    logo_p = para("", indent=False, align=C, space_after=36)
    logo_p.paragraph_format.space_before = Pt(46)
    logo_p.add_run().add_picture(str(LOGO), width=Cm(4.6))
    para("Department of Humanities and Social Sciences", indent=False, align=C, size=12, space_after=1)
    para("Indian Institute of Technology Kharagpur", indent=False, align=C, size=12, space_after=1)
    para("Kharagpur, West Bengal 721302", indent=False, align=C, size=12, space_after=36)
    para("© 2026. All rights reserved", indent=False, align=C, size=11)


# ============================================================================= build
setup_styles()
cover = doc.sections[0]
cover.page_height, cover.page_width = Cm(29.7), Cm(PAGE_W)
cover.left_margin = cover.right_margin = Cm(MARGIN)
cover.top_margin, cover.bottom_margin = Cm(2.0), Cm(2.0)
cover.footer.is_linked_to_previous = False
cover_page()

# ----------------------------------------------------------------------------- contents page
contents = new_section(1, new_page=True)
contents.footer.is_linked_to_previous = False
for title, instr in [("Contents", 'TOC \\o "1-2" \\h \\z \\u'),
                     ("List of Figures", 'TOC \\h \\z \\t "Figure Caption,1"'),
                     ("List of Tables", 'TOC \\h \\z \\t "Table Caption,1"')]:
    para(title, indent=False, size=13, bold=True, space_after=4).paragraph_format.space_before = Pt(8)
    add_field(para("", indent=False), instr)

# ----------------------------------------------------------------------------- paper front matter (spans the page)
front = new_section(1, new_page=True)
pg = OxmlElement("w:pgNumType")
pg.set(qn("w:fmt"), "decimal")
pg.set(qn("w:start"), "1")
front._sectPr.append(pg)
footer_page_number(front)

para(TITLE, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=22, space_after=12)
authors_tbl = doc.add_table(rows=2, cols=len(AUTHORS))
authors_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
for j, name in enumerate(AUTHORS):
    for i, (txt, it) in enumerate([(name, False), ("Dept. of Humanities and Social Sciences\nIIT Kharagpur", True)]):
        p = authors_tbl.rows[i].cells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(txt)
        r.italic = it
        r.font.size = Pt(11 if i == 0 else 9)
para("", indent=False, space_after=6)

single_ame, rich_ame, low_ame = ame("Single"), ame("Rich class"), ame("Lower-middle class")
mid_ame, high_ame, male_ame = ame(MID), ame(HIGH), ame("Male")
cgpa_ame, friend_ame = ame("CGPA"), ame("Has friend circle")

abstract = (
    "Depression is one of the most common mental health problems among university students, yet it is unclear which "
    "personal, social and economic characteristics separate depressed students from others once they are considered "
    f"together. Using survey data on {R['n']} university students aged {R['age_min']}–{R['age_max']} with a CGPA above "
    f"6.00 on a 10-point scale, of whom {R['n_depressed']} ({R['pct_depressed']}%) report depression, this paper "
    "estimates binary logit and probit models with average marginal effects, supported by chi-square and t-tests, "
    f"multicollinearity, goodness-of-fit and robustness checks. Being single is the strongest risk factor "
    f"({pp(single_ame['dy/dx'])}, {fmt_p(single_ame['P>|z|'])}). Compared with students whose monthly living expense "
    f"is low, students with mid ({pp(mid_ame['dy/dx'])}) and high ({pp(high_ame['dy/dx'])}) expenses are more likely "
    f"to be depressed, and male students are {abs(100 * male_ame['dy/dx']):.1f} points less likely to be depressed "
    f"than female students ({fmt_p(male_ame['P>|z|'])}). A rich family background is protective only at the 10% "
    f"level, and CGPA, semester, age, living arrangement and having a friend circle show no significant effect. The "
    f"logit model is jointly significant (LR χ²({R['logit_lr_df']}) = {R['logit_lr_chi2']:.2f}, "
    f"{fmt_p(R['logit_lr_p'])}), well calibrated (Hosmer–Lemeshow p = {R['logit_hl_p']:.3f}) and has an AUC of "
    f"{R['logit_auc']:.3f}; the probit model gives almost identical effects. The results suggest that student mental "
    "health support should prioritise social connection and financial pressure rather than academic performance alone."
)
p = para("", indent=True)
r = p.add_run("Abstract—")
r.bold, r.italic = True, True
add_rich(p, abstract, bold=True, size=9)
p = para("", indent=True, space_after=6)
r = p.add_run("Index Terms—")
r.bold, r.italic = True, True
add_rich(p, "Depression, university students, binary logit, probit, marginal effects, relationship status, "
         "socioeconomic status, mental health.", bold=True, size=9)

# ----------------------------------------------------------------------------- two-column body
new_section(2)

heading("Introduction")
subheading("Background")
para("University life combines several major transitions within a few years: students move away from home, form new "
     "friendships and romantic relationships, manage their own money and face continuous academic evaluation. For many "
     "students these changes are a source of persistent stress, and depression—marked by low mood, loss of interest, "
     "fatigue and feelings of worthlessness—is among the most frequent mental health problems in this group. A "
     "systematic review of university students reported an average depression prevalence of about 30%, well above the "
     "general population [@ibrahim]. Student depression also lowers academic productivity and raises the risk of "
     "dropping out [@hysen; @eisenberg], so knowing which students are most at risk matters to universities deciding "
     "where to direct limited counselling resources.")
para("Theory points to many interrelated causes: stressful life events acting on individual vulnerability [@monroe], "
     "the protective role of social support [@cohen], a socioeconomic gradient in depression [@lorant] and higher "
     "depression among women [@nolen]. This paper brings these strands together and tests them jointly on one "
     "well-defined group of students.")
subheading("Problem Statement")
para("Many studies examine one factor at a time—for example grades or family income—and therefore cannot tell whether "
     "an apparent effect is real or reflects other, correlated characteristics. If single students are also more "
     "likely to live away from family, a simple comparison cannot separate the two. There is also a widespread "
     "assumption that poor grades are a main driver of depression; if this is wrong, support programmes focused on "
     "academic help will miss many students in need. This study estimates multivariate probability models that "
     "isolate the contribution of each factor while holding the others constant.")
subheading("Objectives and Research Questions")
para("The study aims to (i) measure the prevalence of depression in the sample; (ii) compare depression across "
     "gender, family background, relationship status, living style, social support and health; "
     "(iii) estimate the effect of each factor with logit and probit models; (iv) express the effects as changes in "
     "probability; (v) test their robustness; and (vi) derive recommendations for universities. It asks whether "
     "academic performance, relationship status, family background and living costs, gender, and social "
     "support are associated with the probability of depression.")
subheading("Hypotheses")
hyp = T["hypothesis_tests"]
para("Based on the literature in Section II, the following hypotheses are tested:")
bullets([f"**{r.Hypothesis}:** {r.Statement}" for r in hyp.itertuples()])
subheading("Contributions and Organization")
para("The paper contributes evidence on the relative importance of social, economic and academic factors measured on "
     "the same students, reports marginal effects that translate results into probability changes, and offers "
     "practical guidance for student support services. Section II reviews the literature, Section III describes the "
     "data and methodology, Section IV presents the results, Section V discusses them, Section VI covers limitations "
     "and future research, and Section VII concludes. The appendix reports the full model output.")

heading("Literature Review")
subheading("The Diathesis–Stress Model")
para("The diathesis–stress model holds that depression develops when stressful events act on a pre-existing "
     "vulnerability [@monroe]. Students face a dense cluster of stressors—examinations, financial obligations, "
     "relationship changes and living away from home—so the model predicts higher depression among students exposed "
     "to more stressors or with fewer buffers. Beck's cognitive theory adds that negative interpretations of such "
     "events turn stress into depressive symptoms [@beck].")
subheading("Social Support and the Stress-Buffering Hypothesis")
para("Social support protects mental health mainly by buffering the impact of stress: people with close, supportive "
     "relationships cope better with the same adversity [@cohen]. Friendships, partners and family reduce the risk of "
     "depression through emotional support and a sense of belonging [@kawachi]. For students, friend circles, partners "
     "and living with family are the most direct sources of such support (H2, H3, H6).")
subheading("Academic Pressure and Performance")
para("Academic stress is widely cited as a driver of student depression, and depression is associated with lower grade "
     "point averages [@eisenberg; @hysen]. The direction of causality is unclear, however, and links often weaken once "
     "social and economic variables are included. H1 tests whether CGPA has an independent association.")
subheading("Socioeconomic Status and Financial Strain")
para("People with lower socioeconomic status have significantly higher odds of depression, with financial strain a "
     "main channel [@lorant]. Debt and money worries are strongly linked to poor mental health [@richardson]. Family "
     "background shapes both resources and expectations, while a student's monthly living expense is a direct measure "
     "of financial pressure (H4a, H4b, H5).")
subheading("Relationships and Gender")
para("The quality and stability of romantic relationships are closely tied to depressive symptoms [@whitton], and the "
     "link between relationship disruption and depression runs in both directions [@bulloch]. Being single can bring loneliness "
     "in early adulthood (H3). Women experience depression roughly twice as often as men [@nolen; @kessler], a gap "
     "attributed to differences in stress exposure, coping styles and biology (H7).")
subheading("Physical Health and Summary")
para("Chronic physical conditions sharply increase the probability of depression [@moussavi], so self-reported health "
     "is expected to be strongly related to depression. When a health measure partly reflects the same symptoms, "
     "however, it cannot be treated as an independent cause—an issue that proves important here. Overall, the "
     "literature identifies academic, social, economic, demographic and health factors, but rarely estimates them "
     "jointly in one sample; this paper addresses that gap.")

heading("Data and Methodology")
subheading("Data")
para(f"The study uses survey data on {R['n']} university students aged {R['age_min']}–{R['age_max']} "
     f"({R['n_male']} male, {R['n_female']} female). The 18–24 range covers the typical undergraduate age. CGPA is "
     "measured on a 10-point scale, and every student in the sample has a CGPA above 6.00. For each student the data "
     "record gender, "
     "age, family background, relationship status, living style, current semester, CGPA, friend circle, self-reported health, "
     "monthly living expense excluding study costs (recorded in three bands) and depression status. "
     f"{tab_ref()} shows the first five records.")
prev = T["raw_data_preview"].drop(columns="Student_ID").rename(columns={
    "Family_background": "Family background", "Relationship_status": "Relationship",
    "Living_style": "Living", "Current_semester": "Semester", "friend_circle": "Friends",
    "Health_condition": "Health", "Living_expense": "Expense"})
table(prev, "Raw Data (First Five Records)", wide=True, font_size=7, decimals=2,
      col_widths=[1.35, 0.8, 1.85, 2.15, 1.45, 1.45, 1.05, 1.25, 1.45, 1.35, 1.6])
subheading("Data Transformation")
para(f"The data contain {R['missing_total']} missing values and {R['duplicates']} duplicate records. Text answers were "
     "converted to 0/1 indicators (depression, male, living with family, friend circle). Family background was coded "
     "as two dummies (lower-middle class, rich class) with middle class as the base. The very small 'complicated' "
     f"relationship group ({int(rate('Relationship status', 'Complicated')['Students'])} students, none depressed) "
     "would be predicted perfectly, so it was merged and relationship status became **single vs. not single**. "
     "Monthly living expense (excluding study costs) is grouped into three bands cut at the rounded sample thirds: "
     "**low** (below 10,000), **mid** (10,000–14,999) and **high** (15,000 or more); it enters the models as two "
     "dummies (mid, high) with low as the base.")
para(f"Self-reported health is almost perfectly aligned with depression: {R['health_critical_rate']:.0f}% of students "
     f"in critical health and {R['health_good_rate']:.0f}% in good health are depressed. Including it in the main "
     "model would cause quasi-complete separation, and it very likely measures the same condition as the outcome. It "
     f"is therefore analysed descriptively and entered only in one robustness model. {tab_ref()} describes the final "
     "variables.")
dic = T["variable_dictionary"].rename(columns={"Variable": "Variable name"})
table(dic[["Variable name", "Description and coding", "Data type"]], "Description of Variables", wide=True,
      font_size=7.5, col_widths=[2.8, 9.6, 4.6])
subheading("Econometric Models")
para("Because the outcome is binary, ordinary least squares can predict values outside 0–1 and is heteroskedastic by "
     "construction, so binary choice models are used [@wooldridge; @long]. The logit model is")
para("Pr(yᵢ = 1 | xᵢ) = Λ(xᵢ′β) = exp(xᵢ′β) / [1 + exp(xᵢ′β)]          (1)", style="Equation")
para("with xᵢ′β = β₀ + β₁Male + β₂Age + β₃CGPA + β₄Semester + β₅MidExpense + β₆HighExpense + β₇LowerMid + "
     "β₈Rich + β₉Single + β₁₀LiveFamily + β₁₁FriendCircle. Coefficients are estimated by maximum likelihood; exp(β) is "
     "the odds ratio. The probit model replaces Λ with the standard normal CDF Φ:", indent=False)
para("Pr(yᵢ = 1 | xᵢ) = Φ(xᵢ′β)          (2)", style="Equation")
para("Because coefficients of the two models are on different scales, results are reported mainly as average "
     "marginal effects (AME): the average change in Pr(y = 1) when a continuous variable rises by one unit, or when a "
     "0/1 variable changes from 0 to 1 (Stata: margins, dydx(*)).", indent=False)
subheading("Diagnostics, Robustness and Software")
para("Fit is assessed with the likelihood-ratio test, McFadden's pseudo-R², AIC/BIC, classification tables, the area "
     "under the ROC curve and the Hosmer–Lemeshow test [@hosmer]; multicollinearity with variance inflation factors. "
     "Robustness checks use robust standard errors, a linear probability model (LPM), a model adding health status, "
     "and separate models for male and female students. Estimation used Python (pandas, SciPy, statsmodels "
     "[@seabold]); an equivalent Stata do-file reproduces every step [@stata]. The data are anonymised, contain no "
     "personal identifiers, and results are reported only at group level.")

heading("Empirical Results")
subheading("Descriptive Analysis")
s_rate, r_rate, c_rate = (rate("Relationship status", g) for g in ("Single", "In a Relationship", "Complicated"))
mid, low, rich = (rate("Family background", g) for g in ("Middle class", "Lower Middle class", "Rich class"))
bach, fam = rate("Living style", "Bachelor"), rate("Living style", "Family")
fem, mal = rate("Gender", "Female"), rate("Gender", "Male")
fr_y, fr_n = rate("Friend circle", "Yes"), rate("Friend circle", "No")
ex_l, ex_m, ex_h = (rate("Living expense", g) for g in ("Low", "Mid", "High"))
para(f"Depression is common: {R['pct_depressed']}% of students—almost three in ten—report it ({fig_ref()}), in line "
     "with international estimates of around 30% [@ibrahim]. The outcome is balanced enough for a binary choice model, "
     "with enough depressed students relative to the number of regressors for stable estimation.")
figure("fig01_depression_distribution", "Distribution of the dependent variable (depression status).")
para(f"{fig_ref()} compares depression rates across groups. Relationship status shows the largest gap: "
     f"{s_rate['Depression rate (%)']:.1f}% of single students are depressed versus {r_rate['Depression rate (%)']:.1f}% "
     f"of students in a relationship and none of the {int(c_rate['Students'])} in a complicated relationship. "
     f"Middle-class students have the highest rate ({mid['Depression rate (%)']:.1f}%), above lower-middle class "
     f"({low['Depression rate (%)']:.1f}%) and rich students ({rich['Depression rate (%)']:.1f}%). Students living as "
     f"bachelors are more often depressed ({bach['Depression rate (%)']:.1f}%) than those living with family "
     f"({fam['Depression rate (%)']:.1f}%). Students with a low living expense are less often depressed "
     f"({ex_l['Depression rate (%)']:.1f}%) than those with a mid ({ex_m['Depression rate (%)']:.1f}%) or high "
     f"({ex_h['Depression rate (%)']:.1f}%) expense. Gender and friend circle show small raw differences (male "
     f"{mal['Depression rate (%)']:.1f}%, female {fem['Depression rate (%)']:.1f}%; friend circle "
     f"{fr_y['Depression rate (%)']:.1f}% vs. {fr_n['Depression rate (%)']:.1f}%). These raw comparisons ignore "
     "overlaps between characteristics—single students are more often male and less often live with family—which the "
     "regression models separate.")
figure("fig02_rates_by_group", "Depression rate by student characteristics (dashed line: sample average).", wide=True)
para(f"{fig_ref()} shows how closely self-reported health tracks depression: none of the students in good health and "
     f"only {R['health_normal_rate']:.1f}% in normal health are depressed, against {R['health_abnormal_rate']:.1f}% in "
     f"abnormal and {R['health_critical_rate']:.0f}% in critical health. Health therefore nearly predicts the outcome "
     "and is not used as a regressor in the main model.")
figure("fig03_health_condition", "Depression by self-reported health condition.")
cg, se, ag = (row("t_tests", v) for v in ("CGPA", "Current semester", "Age (years)"))
para(f"{fig_ref()} compares the numerical variables. Mean CGPA ({cg['Mean (not depressed)']:.2f} vs. "
     f"{cg['Mean (depressed)']:.2f}), age ({ag['Mean (not depressed)']:.1f} vs. {ag['Mean (depressed)']:.1f}) and "
     f"semester ({se['Mean (not depressed)']:.1f} vs. {se['Mean (depressed)']:.1f}) are almost identical for "
     f"non-depressed and depressed students. Depression rates by age and semester ({fig_ref(2)}) show no clear trend; the very small groups at ages 18 and 19 should be read "
     "with caution.")
figure("fig04_numeric_by_depression", "Distribution of CGPA, age and semester by depression status.", wide=True)
figure("fig05_rate_by_age_semester", "Depression rate by age and semester group.", wide=True)
summ = T["summary_statistics"].set_index("Variable")
para(f"{tab_ref()} reports summary statistics. Students are on average {summ.loc['Age (years)', 'Mean']:.1f} years "
     f"old with a mean CGPA of {summ.loc['CGPA', 'Mean']:.2f} out of 10; {100 * summ.loc[MID, 'Mean']:.0f}% have a mid "
     f"and {100 * summ.loc[HIGH, 'Mean']:.0f}% a high living expense; {100 * summ.loc['Single', 'Mean']:.0f}% are single, {100 * summ.loc['Has friend circle', 'Mean']:.0f}% "
     f"have a friend circle and {100 * summ.loc['Lives with family', 'Mean']:.0f}% live with family.")
table(T["summary_statistics"], "Summary Statistics of Coded Variables", font_size=7, decimals=3)
cm = T["correlation_matrix"].set_index(T["correlation_matrix"].columns[0])
para(f"The correlation matrix ({fig_ref()}) shows modest correlations among the regressors; the largest are between "
     f"the two expense dummies ({cm.loc[MID, HIGH]:.2f}, by construction), between male and living with family "
     f"({cm.loc['Male', 'Lives with family']:.2f}) and between high expense and lower-middle class "
     f"({cm.loc[HIGH, 'Lower-middle class']:.2f}). The very high correlation between depression and health "
     f"level ({cm.loc['Depression', 'Health level (1-4)']:.2f}) confirms the separation issue.".replace("-0.", "−0."))
figure("fig06_correlation_matrix", "Correlation matrix of coded variables.", wide=True, width_cm=12.5)

subheading("Bivariate Tests")
ch = T["chi_square_tests"].set_index("Variable")
para(f"{tab_ref()} reports chi-square tests of independence. Relationship status (χ² = "
     f"{ch.loc['Relationship status', 'Chi-square']:.2f}, {fmt_p(ch.loc['Relationship status', 'p-value'])}, Cramér's "
     f"V = {ch.loc['Relationship status', CRAMER]:.2f}) and family background (χ² = "
     f"{ch.loc['Family background', 'Chi-square']:.2f}, {fmt_p(ch.loc['Family background', 'p-value'])}) are "
     f"significantly associated with depression, as are living style ({fmt_p(ch.loc['Living style', 'p-value'])}) and "
     f"health (V = {ch.loc['Health condition', CRAMER]:.2f}). Living expense is also associated with depression "
     f"(χ² = {ch.loc['Living expense', 'Chi-square']:.2f}, {fmt_p(ch.loc['Living expense', 'p-value'])}); gender and "
     f"friend circle are not significant. {tab_ref(2)} shows that none of the numerical variables differs "
     f"significantly: CGPA ({fmt_p(cg['t-test p-value'])}), age ({fmt_p(ag['t-test p-value'])}) and semester "
     f"({fmt_p(se['t-test p-value'])}).")
table(T["chi_square_tests"], "Chi-Square Tests of Independence with Depression", wide=True, font_size=7.5)
table(T["t_tests"], "Difference in Means Between Depressed and Non-Depressed Students", wide=True, font_size=7.5)

subheading("Multicollinearity")
vif_tab = T["vif"].set_index("Variable")
para(f"The largest variance inflation factor is {R['vif_max']:.2f} ({tab_ref()}), below the common threshold of 10 and "
     "the stricter 5, so multicollinearity is not a concern. The highest values belong to the two living-expense "
     f"dummies ({vif_tab.loc[MID, 'VIF']:.2f} and {vif_tab.loc[HIGH, 'VIF']:.2f}), as expected for dummies of the same "
     "variable that also overlap with family background.")
table(T["vif"], "Variance Inflation Factors", font_size=7.5)

subheading("Binary Logit Model")
lt = T["logit_coefficients"].copy()
lt["Sig."] = lt["Sig."].fillna("")
para(f"{tab_ref()} presents the logit estimates. The model is highly significant (LR χ²({R['logit_lr_df']}) = "
     f"{R['logit_lr_chi2']:.2f}, {fmt_p(R['logit_lr_p'])}) with McFadden's pseudo-R² = {R['logit_pseudo_r2']:.3f}, a "
     "reasonable value for individual cross-sectional data.")
table(lt, f"Binary Logit Estimates (N = {R['n']}; Dependent Variable: Depression)", wide=True, font_size=7.5, decimals=4)
c_s, c_mid, c_high, c_r, c_male = (coef(v) for v in ("Single", MID, HIGH, "Rich class", "Male"))
para(f"Holding other factors constant, being single has the strongest effect (β = {c_s['Coefficient']:.3f}, "
     f"{fmt_p(c_s['P>|z|'])}). Mid (β = {c_mid['Coefficient']:.3f}, {fmt_p(c_mid['P>|z|'])}) and high living expense "
     f"(β = {c_high['Coefficient']:.3f}, {fmt_p(c_high['P>|z|'])}) raise the probability relative to low expense, and "
     f"male students are less likely to be depressed (β = {c_male['Coefficient']:.3f}, {fmt_p(c_male['P>|z|'])}). The "
     f"rich-family coefficient is negative but not significant (β = {c_r['Coefficient']:.3f}, {fmt_p(c_r['P>|z|'])}). "
     "CGPA, age, semester, lower-middle class, living with family and friend circle are not significant.")
o_s, o_mid, o_high, o_male = (odds(v) for v in ("Single", MID, HIGH, "Male"))
para(f"In odds-ratio terms ({tab_ref()}, {fig_ref()}), single students have {o_s['Odds ratio']:.1f} times the odds "
     f"of depression of students in a relationship (95% CI {o_s['95% CI lower']:.1f}–{o_s['95% CI upper']:.1f}); "
     f"students with a mid or high living expense have {o_mid['Odds ratio']:.1f} and {o_high['Odds ratio']:.1f} times "
     f"the odds of those with a low expense; and male students have {100 * (1 - o_male['Odds ratio']):.0f}% lower odds "
     "than female students. The wide confidence intervals of the expense odds ratios reflect the modest size of each "
     "band.")
ot = T["logit_odds_ratios"].copy()
ot["Sig."] = ot["Sig."].fillna("")
table(ot, "Odds Ratios from the Logit Model", font_size=7, decimals=3)
figure("fig07_logit_odds_ratios", "Logit odds ratios with 95% confidence intervals (log scale).", wide=True, width_cm=11.5)

subheading("Marginal Effects")
la = T["logit_marginal_effects_ame"].copy()
la["Sig."] = la["Sig."].fillna("")
para(f"{tab_ref()} and {fig_ref()} report the average marginal effects. For a student with average characteristics "
     f"the predicted probability of depression is {100 * R['logit_pred_at_means']:.1f}%.")
table(la, "Average Marginal Effects after Logit (dy/dx)", wide=True, font_size=7.5, decimals=4)
figure("fig08_logit_marginal_effects", "Average marginal effects after logit with 95% confidence intervals.", wide=True, width_cm=11.5)
bullets([
    f"Being **single** raises the probability of depression by {100 * single_ame['dy/dx']:.1f} percentage points "
    f"({fmt_p(single_ame['P>|z|'])})—by far the largest effect.",
    f"Compared with a low living expense, a **mid expense** raises it by {100 * mid_ame['dy/dx']:.1f} points "
    f"({fmt_p(mid_ame['P>|z|'])}) and a **high expense** by {100 * high_ame['dy/dx']:.1f} points "
    f"({fmt_p(high_ame['P>|z|'])}). The two effects are almost equal, so the step is between low and any higher "
    "expense rather than a steady rise.",
    f"**Male** students are {abs(100 * male_ame['dy/dx']):.1f} points less likely to be depressed "
    f"({fmt_p(male_ame['P>|z|'])}).",
    f"A **rich family** lowers the probability by {abs(100 * rich_ame['dy/dx']):.1f} points relative to middle "
    f"class, significant only at 10% ({fmt_p(rich_ame['P>|z|'])}); the **lower-middle class** effect is close to "
    f"zero ({pp(low_ame['dy/dx'])}, {fmt_p(low_ame['P>|z|'])}).",
    f"**CGPA** ({100 * cgpa_ame['dy/dx']:+.1f} points per grade point on the 10-point scale, "
    f"{fmt_p(cgpa_ame['P>|z|'])}) and **friend circle** ({fmt_p(friend_ame['P>|z|'])}) have no significant effect.",
])
para("Marginal effects at the means (Stata's older mfx) lead to the same conclusions (see Appendix).")

subheading("Binary Probit Model")
pt = T["probit_coefficients"].copy()
pt["Sig."] = pt["Sig."].fillna("")
pa = T["probit_marginal_effects_ame"].copy()
pa["Sig."] = pa["Sig."].fillna("")
ps, pmid, phigh, pm = (ame(v, "probit") for v in ("Single", MID, HIGH, "Male"))
para(f"The probit model ({tab_ref()}) is equally significant (LR χ²({R['logit_lr_df']}) = {R['probit_lr_chi2']:.2f}, "
     f"pseudo-R² = {R['probit_pseudo_r2']:.3f}) with the same signs and significance pattern: single, mid and high "
     f"expense and male are significant at 5%, rich class is not ({fmt_p(coef('Rich class', 'probit')['P>|z|'])}). "
     f"Its marginal effects ({tab_ref(2)}) are almost identical to the logit ones: single {pp(ps['dy/dx'])} (logit "
     f"{100 * single_ame['dy/dx']:.1f}), mid expense {pp(pmid['dy/dx'])}, high expense {pp(phigh['dy/dx'])}, male "
     f"{pp(pm['dy/dx'])}. The choice of link function makes no practical difference.")
table(pt, f"Binary Probit Estimates (N = {R['n']}; Dependent Variable: Depression)", wide=True, font_size=7.5, decimals=4)
table(pa, "Average Marginal Effects after Probit (dy/dx)", wide=True, font_size=7.5, decimals=4)

subheading("Goodness of Fit")
fs = T["model_fit_statistics"].set_index("Model").T.reset_index().rename(columns={"index": "Statistic"})
cl = T["classification_tables"].copy()
probit_hl_p = T["model_fit_statistics"].set_index("Model").loc["Probit", "HL p-value"]
para(f"Both models have almost the same likelihood, AIC and AUC ({tab_ref()}). The Hosmer–Lemeshow test does not reject "
     f"the logit model (χ² = {R['logit_hl_chi2']:.2f}, p = {R['logit_hl_p']:.3f}), so its predicted and observed rates "
     f"agree across risk groups; for the probit model the test is rejected (p = {probit_hl_p:.3f}), which is one reason "
     "the logit is kept as the main model. At a 0.5 cut-off the logit model classifies {100 * R['logit_correct_05']:.1f}% of students "
     f"correctly, with specificity {100 * R['logit_spec_05']:.1f}% but sensitivity only {100 * R['logit_sens_05']:.1f}% "
     f"({tab_ref(2)}). Using the sample share of depressed students ({R['pct_depressed'] / 100:.3f}) as cut-off "
     f"identifies {100 * cl.loc[1, 'Sensitivity']:.1f}% of depressed students while correctly classifying "
     f"{100 * cl.loc[1, 'Specificity']:.1f}% of the others.")
table(fs, "Goodness-of-Fit Statistics", font_size=7, decimals=3)
table(cl, "Classification Tables (Cut-off 0.5 and Sample Share)", wide=True, font_size=7.5, decimals=3)
prof = T["predicted_probability_profiles"]
prof_s, prof_n = (prof.loc[prof["Single"] == s, "Predicted probability"] for s in (1, 0))
para(f"The ROC curves ({fig_ref()}) give an AUC of {R['logit_auc']:.3f} (logit) and {R['probit_auc']:.3f} (probit): a "
     f"randomly chosen depressed student has a {100 * R['logit_auc']:.0f}% chance of a higher predicted probability than "
     f"a randomly chosen non-depressed student. {fig_ref(2)} shows the size of the relationship effect: with other "
     f"variables at their means, the predicted probability for single students ranges from {100 * prof_s.min():.0f}% "
     f"to {100 * prof_s.max():.0f}% across the CGPA range, versus {100 * prof_n.min():.0f}–{100 * prof_n.max():.0f}% "
     "for other students. The downward slope of both lines is not statistically significant.")
figure("fig09_roc_curves", "ROC curves of the logit and probit models.")
figure("fig10_predicted_probability_cgpa", "Predicted probability of depression by CGPA and relationship status.")

subheading("Hypothesis Tests")
sup = hyp[hyp["Result"] == "Supported"]["Hypothesis"].tolist()
weak = hyp[hyp["Result"].str.startswith("Weakly")]["Hypothesis"].tolist()
nots = hyp[hyp["Result"].str.startswith("Not")]["Hypothesis"].tolist()
assert sup == ["H3", "H5", "H7"] and not weak, "text below describes H3, H5 and H7 as the supported hypotheses"
para(f"{tab_ref()} summarises the hypothesis tests (logit model; supported if the expected sign holds with p < 0.05). "
     "H3, H5 and H7 are supported: single students, students with a high living expense and female students are more "
     f"likely to be depressed. {', '.join(nots[:-1])} and {nots[-1]} are not supported: academic performance, friend "
     "circle, family background and living with family have no significant independent effect at 5%, although the "
     "rich-family marginal effect is significant at 10%.")
table(hyp[["Hypothesis", "Variable", "Expected sign", "Logit coefficient", "AME", "P>|z|", "Result"]],
      "Summary of Hypothesis Tests", wide=True, font_size=7.5, decimals=4)

subheading("Robustness Checks")
rb = T["robustness_models"].fillna("")
rb = rb.rename(columns=lambda c: c.replace("(", "").replace(")", "").replace(" robust SE", " rob.")
               .replace("Logit + poor health", "Logit+health"))
para(f"{tab_ref()} compares the main model with six alternatives, and {fig_ref()} compares the marginal effects of "
     "the logit, probit and linear probability models.")
table(rb, "Robustness Checks Across Model Specifications", wide=True, font_size=7,
      col_widths=[2.8, 2.0, 2.0, 2.0, 2.0, 2.1, 2.0, 2.0])
note("Significance levels: *** p < 0.01, ** p < 0.05, * p < 0.10. Standard errors in parentheses. "
     "Models 6–7 are linear probability models with robust standard errors.")
figure("fig11_robustness_comparison", "Marginal effects across logit, probit and linear probability models.", wide=True, width_cm=11.5)
sub = T["subsample_marginal_effects"]


def sub_me(sample, var):
    r = sub[(sub["Sample"] == sample) & (sub["Variable"] == var)].iloc[0]
    return r["dy/dx"], r["P>|z|"]


m_s, f_s = sub_me("Male students", "Single"), sub_me("Female students", "Single")
m_r = sub_me("Male students", "Rich class")
m_mid, f_mid = sub_me("Male students", MID), sub_me("Female students", MID)
m_high, f_high = sub_me("Male students", HIGH), sub_me("Female students", HIGH)
f_g = sub_me("Female students", "CGPA")
cmp_ = T["marginal_effects_comparison"].set_index("Variable")
bullets([
    "**Robust standard errors** change the standard errors only slightly; every variable keeps its sign and its "
    "significance at the 5% level.",
    f"**The LPM** (R² = {R['lpm_r2']:.3f}) gives almost the same effects for single "
    f"({pp(cmp_.loc['Single', 'LPM coefficient'])}) and male ({pp(cmp_.loc['Male', 'LPM coefficient'])}). Its "
    f"living-expense effects are smaller ({pp(cmp_.loc[MID, 'LPM coefficient'])} and "
    f"{pp(cmp_.loc[HIGH, 'LPM coefficient'])}) and significant only at 10% (p = {cmp_.loc[MID, 'LPM p (robust)']:.3f} "
    f"and {cmp_.loc[HIGH, 'LPM p (robust)']:.3f}), so the expense result is somewhat less robust than the others.",
    f"**Adding poor health** raises the pseudo-R² from {R['logit_pseudo_r2']:.3f} to {R['health_model_pseudo_r2']:.3f} "
    f"and the AUC to {R['health_model_auc']:.3f} (LR χ²(1) = {R['lr_health_chi2']:.1f}, {fmt_p(R['lr_health_p'])}); "
    f"the software flags quasi-complete separation (a fraction {QS_FRACTION} of observations predicted perfectly), "
    "confirming that health largely measures the outcome itself. Single and male remain significant; the "
    "living-expense and rich-class coefficients lose significance.",
    f"**By gender,** a logit cannot be estimated for men because none of the {R['male_not_single_n']} non-single male "
    f"students is depressed, so LPMs are used. Being single is a strong risk factor for men ({pp(m_s[0])}, "
    f"{fmt_p(m_s[1])}) and women ({pp(f_s[0])}, {fmt_p(f_s[1])}). Among men, mid ({pp(m_mid[0])}, "
    f"{fmt_p(m_mid[1])}) and high living expense ({pp(m_high[0])}, {fmt_p(m_high[1])}) raise the risk and a rich "
    f"family lowers it ({pp(m_r[0])}, {fmt_p(m_r[1])}); among women the expense effects are smaller and not "
    f"significant (p = {f_mid[1]:.2f}, {f_high[1]:.2f}), while a higher CGPA is associated with lower risk "
    f"({100 * f_g[0]:+.1f} points per grade point, {fmt_p(f_g[1])}). With fewer than 50 depressed students per "
    "subsample, these differences are indicative.",
])

heading("Discussion")
subheading("Key Findings")
para(f"**Relationships matter more than grades.** Single students are {100 * single_ame['dy/dx']:.0f} percentage "
     "points more likely to be depressed, an effect that is large, highly significant and stable across every "
     "specification and for both men and women. This strongly supports the stress-buffering view of close "
     "relationships [@cohen].")
para(f"**Financial pressure.** Students whose monthly living expense is mid or high are about "
     f"{100 * (mid_ame['dy/dx'] + high_ame['dy/dx']) / 2:.0f} points more likely to be depressed than students with a "
     "low expense, after controlling for family background. The effect appears as a step from low to any higher "
     "expense rather than a steady rise, consistent with research linking money worries to poor mental health "
     "[@richardson]. Family background itself matters less once expenses are controlled: rich students are less "
     "likely to be depressed, but only at the 10% level, and middle-class students—who have the highest raw rate—may "
     "face high costs without the cushion of richer families [@lorant].")
para(f"**The gender gap.** After controls, female students are about {abs(100 * male_ame['dy/dx']):.0f} points more "
     "likely to be depressed. The gap is hidden in the raw data—where men show a slightly higher rate—because male "
     "students are more often single. The adjusted result matches the established gender difference [@nolen; @kessler].")
para("**What did not matter.** CGPA, semester and age have no significant relationship with depression, and neither "
     "do a friend circle or living with family. Because every student in the sample has a CGPA above 6.00, the result "
     "says that grades do not matter much within this range, not that very low grades are harmless. Living with "
     "family looks protective in the raw data, but the effect disappears once relationship status and gender are "
     "controlled. The support that matters most here seems to be an intimate relationship rather than a wider "
     "friend group.")
subheading("Implications for Universities")
para("Since grades do not predict depression in this sample, services focused mainly on academic difficulty will miss "
     "many depressed students. Universities should (i) invest in programmes that reduce loneliness and build close "
     "connections—peer mentoring, small-group and residential community activities; (ii) combine mental health support "
     "with financial guidance—budgeting help, affordable housing and meals—that reaches students with higher living "
     "costs, including middle-income students; (iii) design gender-sensitive outreach, since women face higher adjusted "
     "risk while men may be less likely to seek help; and (iv) "
     f"use simple risk scores to prioritise screening—the logit model identifies {100 * cl.loc[1, 'Sensitivity']:.0f}% of "
     "depressed students at the sample-share cut-off—without replacing clinical assessment.")
subheading("Theoretical Implications")
para("The results support the diathesis–stress and stress-buffering frameworks [@monroe; @cohen]: financial pressure "
     "raises the probability of depression while close relationships protect against it. They suggest that for young "
     "adults the protective effect comes from intimate relationships more than from the size of the social network, "
     "and the step pattern of living expenses suggests that moving from a low to a higher cost of living—rather than "
     "the exact amount—is what adds pressure.")

heading("Limitations and Future Research")
para("The data are cross-sectional, so the estimates are associations rather than causal effects; depression may "
     "itself affect relationship status. Depression and health are self-reported indicators rather than clinical "
     f"diagnoses or validated scales. With {R['n']} students and {R['n_depressed']} depressed cases, small categories "
     "had to be merged and the gender subsample results are indicative. Living expense is measured in three bands, "
     "which loses detail within each band, and the expense effects are weaker in the linear probability model. The "
     "sample covers only students with a CGPA above 6.00 and is not representative of all university students, which "
     "limits generalisation; factors such as family history, sleep, physical activity and substance "
     "use are not observed. Future work should follow students over several semesters, use validated severity "
     "scales such as the PHQ-9 with ordered models, draw larger and more diverse samples covering the full grade range, "
     "and measure the quality—not only the presence—of social support.")

heading("Conclusion")
para("Depression affects almost three in ten students in this sample. It is driven less by grades and more by "
     "students' relationships and financial circumstances: being single is the strongest risk factor, a mid or high "
     "living expense raises the risk compared with a low one, and male students are less likely to be depressed than "
     "female students; a rich family background is protective only weakly. These findings are consistent across the "
     "logit and probit models, and the relationship and gender effects also hold in the linear probability model. "
     "Protecting student mental health therefore requires attention to social connection and financial wellbeing, "
     "not only to academic support.")

# ----------------------------------------------------------------------------- references
doc.add_heading("References", level=1)
for i, key in enumerate(CITED, 1):
    p = para("", indent=False)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    p.paragraph_format.space_after = Pt(1)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(f"[{i}] {REFERENCES[key]}").font.size = Pt(8)
unused = set(REFERENCES) - set(CITED)
assert not unused, f"references never cited: {unused}"

# ----------------------------------------------------------------------------- appendix (single column, continuous)
new_section(1)
doc.add_heading("Appendix: Full Model Output", level=1)
para("Full estimation output from the analysis notebook (notebooks/student_depression_analysis.ipynb, Python "
     "statsmodels). The Stata do-file stata/student_depression_analysis.do contains the equivalent commands (summarize, "
     "tabulate, logit, probit, margins, estat classification, estat gof, lroc, regress). All tables are also in "
     "results/Student_Depression_Analysis_Results.xlsx.", indent=False, space_after=4)
for title, fname in [
    ("A. Binary logit model", "logit_summary"),
    ("B. Binary probit model", "probit_summary"),
    ("C. Logit model with robust standard errors", "logit_robust_summary"),
    ("D. Linear probability model (robust standard errors)", "lpm_summary"),
    ("E. Logit model including poor health", "logit_health_summary"),
    ("F. Linear probability model: male students", "lpm_male_summary"),
    ("G. Linear probability model: female students", "lpm_female_summary"),
]:
    p = para(title, indent=False, italic=True, space_after=2)
    p.paragraph_format.keep_with_next = True
    code_block((RES / "model_output" / f"{fname}.txt").read_text())
mem = T["logit_marginal_effects_mem"].copy()
mem["Sig."] = mem["Sig."].fillna("")
table(mem, "Marginal Effects at the Means after Logit", font_size=7.5, decimals=4)
sp = sub.pivot(index="Variable", columns="Sample", values=["dy/dx", "P>|z|"])
sp.columns = [f"{b}: {a}" for a, b in sp.columns]
table(sp.reset_index(), "Marginal Effects for Male and Female Students (LPM, Robust SE)", font_size=7.5, decimals=4)

doc.core_properties.title = TITLE
doc.core_properties.author = ", ".join(AUTHORS)
doc.core_properties.comments = ""
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print("Saved", OUT.relative_to(ROOT))
